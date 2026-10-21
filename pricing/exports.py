"""Excel 导出：核价单与核价记录清单。"""

from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

THIN = Side(style="thin", color="D8DEE6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="123F55")
LIGHT_FILL = PatternFill("solid", fgColor="F2F6F8")
TITLE_FONT = Font(name="微软雅黑", size=16, bold=True, color="123F55")
HEADER_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10, color="28323C")
MUTED_FONT = Font(name="微软雅黑", size=9, color="6B7785")


def _style_cell(cell, *, font=BODY_FONT, fill=None, align="left", border=True):
    cell.font = font
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    if fill:
        cell.fill = fill
    if border:
        cell.border = BORDER


def build_calc_workbook(calc):
    wb = Workbook()
    ws = wb.active
    ws.title = "核价单"
    widths = [18, 14, 12, 8, 14, 16, 34]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    ws.merge_cells("A1:G1")
    ws["A1"] = "制造业采购成本核价单"
    _style_cell(ws["A1"], font=TITLE_FONT, align="center", border=False)
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:G2")
    ws["A2"] = f"核价编号：{calc.calc_no}    核价日期：{calc.created_at:%Y-%m-%d %H:%M}    核价人：{calc.created_by.get_full_name() or calc.created_by.username}"
    _style_cell(ws["A2"], font=MUTED_FONT, align="center", border=False)

    meta = [
        ("零件编号", calc.part.code, "零件名称", calc.part.name, "材料", calc.part.material.name, "规格", calc.part.specification or "-"),
        ("成品净重(kg)", f"{calc.net_weight:.4f}", "损耗率", f"{calc.loss_rate:.2f}%", "材料单价(元/kg)", f"{calc.material_price:.4f}", "价格日期", str(calc.material_price_date)),
        ("材料投入量(kg)", f"{calc.material_input_weight:.4f}", "规则版本", calc.rule_version, "状态", calc.status_label, "", ""),
    ]
    row = 4
    for meta_row in meta:
        col = 1
        for value in meta_row:
            cell = ws.cell(row=row, column=col, value=value)
            is_label = col % 2 == 1
            _style_cell(cell, font=MUTED_FONT if is_label else BODY_FONT, fill=LIGHT_FILL if is_label else None)
            if is_label and col < 7:
                ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
                col += 1
            col += 1
        row += 1

    row = 9
    headers = ["费用类型", "费用项目", "数量", "单位", "单价(元)", "金额(元)", "计算说明"]
    for index, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=index, value=header)
        _style_cell(cell, font=HEADER_FONT, fill=HEADER_FILL, align="center")
    ws.row_dimensions[row].height = 24

    for detail in calc.details.all():
        row += 1
        values = [
            detail.get_item_type_display(), detail.item_name, float(detail.quantity),
            detail.unit, float(detail.unit_price), float(detail.amount), detail.formula_text,
        ]
        for index, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=index, value=value)
            align = "right" if index in (3, 5, 6) else ("center" if index == 4 else "left")
            _style_cell(cell, align=align)
            if index in (5, 6):
                cell.number_format = "#,##0.00"

    row += 1
    summary_rows = [
        ("理论单位成本", float(calc.unit_cost)),
        ("合理加成率", f"{calc.markup_rate:.2f}%"),
        ("核价参考报价(不含税)", float(calc.reference_price)),
        ("参考报价(含税13%)", float(calc.reference_price) * 1.13),
    ]
    for label, value in summary_rows:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        label_cell = ws.cell(row=row, column=1, value=label)
        _style_cell(label_cell, font=Font(name="微软雅黑", size=10, bold=True), fill=LIGHT_FILL, align="right")
        value_cell = ws.cell(row=row, column=6, value=value)
        _style_cell(value_cell, font=Font(name="微软雅黑", size=10, bold=True, color="123F55"), align="right")
        if isinstance(value, float):
            value_cell.number_format = "#,##0.00"
        row += 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
    ws.cell(row=row, column=1, value=f"核价说明：{calc.remark or '无'}")
    _style_cell(ws.cell(row=row, column=1), font=MUTED_FONT, border=False)

    sheet = wb.create_sheet("报价对比")
    for index, width in enumerate([20, 14, 10, 10, 14, 12, 12, 16, 12], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.merge_cells("A1:I1")
    sheet["A1"] = f"供应商报价对比 · {calc.part.name}（参考报价 {calc.reference_price:.2f} 元/件）"
    _style_cell(sheet["A1"], font=TITLE_FONT, align="center", border=False)

    quote_headers = ["供应商", "报价日期", "原始报价", "单位", "统一口径(元/件)", "偏差金额", "偏差率", "含税口径", "参考区间"]
    for index, header in enumerate(quote_headers, start=1):
        cell = sheet.cell(row=3, column=index, value=header)
        _style_cell(cell, font=HEADER_FONT, fill=HEADER_FILL, align="center")

    from .services import quote_metrics

    row = 3
    for quote in calc.part.quotes.select_related("supplier").order_by("normalized_price"):
        row += 1
        metric = quote_metrics(quote, calc)
        values = [
            quote.supplier.name, str(quote.quote_date), float(quote.unit_price), quote.price_unit,
            float(quote.normalized_price), float(metric["deviation"] or 0),
            float(metric["rate"] or 0) / 100, quote.tax_label,
            f"{calc.deviation_low:.0f}% ~ {calc.deviation_high:.0f}%",
        ]
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=index, value=value)
            _style_cell(cell, align="right" if index in (5, 6, 7) else "left")
            if index == 5:
                cell.number_format = "#,##0.00"
            if index == 6:
                cell.number_format = "+#,##0.00;-#,##0.00"
            if index == 7:
                cell.number_format = "+0.00%;-0.00%"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{calc.calc_no}.xlsx"'
    return response


def build_calc_list_workbook(calculations):
    wb = Workbook()
    ws = wb.active
    ws.title = "核价记录"
    widths = [16, 13, 22, 18, 12, 12, 10, 14, 14, 16, 12]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=11)
    ws["A1"] = "采购成本核价记录清单"
    _style_cell(ws["A1"], font=TITLE_FONT, align="center", border=False)
    ws.row_dimensions[1].height = 30

    headers = [
        "核价编号", "核价日期", "零件名称", "零件编号", "成品净重(kg)", "材料单价",
        "材料成本", "理论单位成本", "参考报价", "核价人", "状态",
    ]
    for index, header in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=index, value=header)
        _style_cell(cell, font=HEADER_FONT, fill=HEADER_FILL, align="center")

    for row_index, calc in enumerate(calculations, start=4):
        values = [
            calc.calc_no, calc.created_at.strftime("%Y-%m-%d"), calc.part.name, calc.part.code,
            float(calc.net_weight), float(calc.material_price), float(calc.material_cost),
            float(calc.unit_cost), float(calc.reference_price),
            calc.created_by.get_full_name() or calc.created_by.username, calc.status_label,
        ]
        for index, value in enumerate(values, start=1):
            cell = ws.cell(row=row_index, column=index, value=value)
            _style_cell(cell, align="right" if index in (5, 6, 7, 8, 9) else "left")
            if index in (6, 7, 8, 9):
                cell.number_format = "#,##0.00"
            if index == 5:
                cell.number_format = "0.0000"

    ws.freeze_panes = "A4"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="cost_calculations.xlsx"'
    return response
